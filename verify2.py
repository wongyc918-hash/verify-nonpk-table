import pandas as pd
from sqlalchemy import create_engine
import concurrent.futures
from openpyxl.styles import PatternFill, Font

# --- CONFIGURATION ---
ORACLE_CONFIG = "oracle+oracledb://matt:oracle123@localhost:22/?service_name=bcauat2"
POLARDB_CONFIG = "postgresql://matt2:Abcd1234@localhost:23/bcauat2-db-team"
INPUT_EXCEL = "tables_to_verify.xlsx"
OUTPUT_EXCEL = "verification_results.xlsx"
MAX_WORKERS = 4  # Adjust based on your CPU/Network capacity

def get_engine(connection_string):
    return create_engine(connection_string)

def normalize_col(series):
    """Normalize a column to string for cross-DB type-safe comparison."""
    if pd.api.types.is_datetime64_any_dtype(series):
        return series.dt.strftime('%Y-%m-%d %H:%M:%S').fillna('')
    elif pd.api.types.is_numeric_dtype(series):
        # Normalize int/float differences (e.g. Oracle int64 vs PolarDB float64 for NULLable cols)
        return series.apply(lambda x: f"{round(float(x), 6)}" if pd.notna(x) else '')
    else:
        # String: strip whitespace and normalize null representations
        return series.astype(str).str.strip().replace({'None': '', 'nan': '', 'NaT': '', '<NA>': ''})


def build_diff_detail(df_ora, df_plr):
    """
    Build a readable side-by-side diff DataFrame.
    - Same row count: shows Row #, Diff Columns, then col [Oracle] / col [PolarDB] side-by-side.
    - Different row count: shows all unmatched rows tagged with Source DB.
    """
    if len(df_ora) != len(df_plr):
        # Row count mismatch — tag each unique row with its source
        df_o = df_ora.copy()
        df_o.insert(0, 'Source DB', 'Oracle')
        df_p = df_plr.copy()
        df_p.insert(0, 'Source DB', 'PolarDB')
        combined = pd.concat([df_o, df_p])
        return combined.drop_duplicates(subset=list(df_ora.columns), keep=False).reset_index(drop=True)

    # Same row count — positional row-by-row comparison
    records = []
    for i in range(len(df_ora)):
        row_o = df_ora.iloc[i]
        row_p = df_plr.iloc[i]
        diff_cols = [c for c in df_ora.columns if row_o[c] != row_p[c]]
        if not diff_cols:
            continue
        record = {'Row #': i + 1, 'Diff Columns': ', '.join(diff_cols)}
        for col in df_ora.columns:
            record[f'{col} [Oracle]'] = row_o[col]
            record[f'{col} [PolarDB]'] = row_p[col]
        records.append(record)
    return pd.DataFrame(records) if records else pd.DataFrame()


def apply_diff_sheet_formatting(ws, df):
    """Highlight header row and color Oracle/PolarDB value columns."""
    HEADER_FILL  = PatternFill("solid", fgColor="4472C4")   # blue header
    ORACLE_FILL  = PatternFill("solid", fgColor="FFF2CC")   # yellow = Oracle
    POLAR_FILL   = PatternFill("solid", fgColor="DDEBF7")   # light blue = PolarDB
    MISMATCH_FILL = PatternFill("solid", fgColor="FCE4D6")  # orange = mismatched value

    header_font = Font(bold=True, color="FFFFFF")

    # Style header row
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = header_font

    # Identify Oracle and PolarDB column indices (1-based)
    headers = [cell.value for cell in ws[1]]
    oracle_indices = {i + 1 for i, h in enumerate(headers) if h and '[Oracle]' in str(h)}
    polar_indices  = {i + 1 for i, h in enumerate(headers) if h and '[PolarDB]' in str(h)}

    # For same-row-count mode: find paired columns and highlight differing cells
    col_pairs = {}  # col_name -> (oracle_col_idx, polardb_col_idx)
    for i, h in enumerate(headers):
        if h and '[Oracle]' in str(h):
            col_name = h.replace(' [Oracle]', '')
            polar_match = next((j + 1 for j, ph in enumerate(headers) if ph == f'{col_name} [PolarDB]'), None)
            if polar_match:
                col_pairs[col_name] = (i + 1, polar_match)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column in oracle_indices:
                cell.fill = ORACLE_FILL
            elif cell.column in polar_indices:
                cell.fill = POLAR_FILL

        # Highlight differing Oracle/PolarDB cell pairs in orange
        for col_name, (ora_idx, plr_idx) in col_pairs.items():
            ora_cell = row[ora_idx - 1]
            plr_cell = row[plr_idx - 1]
            if str(ora_cell.value) != str(plr_cell.value):
                ora_cell.fill = MISMATCH_FILL
                plr_cell.fill = MISMATCH_FILL


def verify_table(table_name):
    print(f"[#] Starting verification for: {table_name}")
    try:
        ora_engine = get_engine(ORACLE_CONFIG)
        plr_engine = get_engine(POLARDB_CONFIG)

        # 1. Fetch data from both sources
        query = f"SELECT * FROM {table_name}"
        df_ora = pd.read_sql(query, ora_engine)
        df_plr = pd.read_sql(query, plr_engine)

        # 2. Normalize column names (Oracle=UPPER, PolarDB=lower) and sort alphabetically
        df_ora.columns = df_ora.columns.str.lower()
        df_plr.columns = df_plr.columns.str.lower()
        df_ora = df_ora.reindex(sorted(df_ora.columns), axis=1)
        df_plr = df_plr.reindex(sorted(df_plr.columns), axis=1)

        # 3. Normalize all values to strings — eliminates type mismatch false positives
        for col in df_ora.columns:
            if col in df_plr.columns:
                df_ora[col] = normalize_col(df_ora[col])
                df_plr[col] = normalize_col(df_plr[col])

        # 4. Sort rows for order-independent comparison
        sort_cols = list(df_ora.columns)
        df_ora = df_ora.sort_values(by=sort_cols, na_position='last').reset_index(drop=True)
        df_plr = df_plr.sort_values(by=sort_cols, na_position='last').reset_index(drop=True)

        # 5. Quick equality check
        if df_ora.equals(df_plr):
            return {"table": table_name, "status": "MATCH", "diff_count": 0, "data": None, "row_ora": len(df_ora), "row_plr": len(df_plr)}

        # 6. Build detailed side-by-side diff
        diff_detail = build_diff_detail(df_ora, df_plr)
        return {
            "table": table_name,
            "status": "MISMATCH",
            "diff_count": len(diff_detail),
            "data": diff_detail,
            "row_ora": len(df_ora),
            "row_plr": len(df_plr)
        }

    except Exception as e:
        return {"table": table_name, "status": "ERROR", "error": str(e), "data": None, "row_ora": None, "row_plr": None}

def main():
    # Load table list from Excel
    input_df = pd.read_excel(INPUT_EXCEL)
    table_list = input_df['table_name'].tolist()

    results_summary = []
    mismatch_details = {}

    # 4. Parallel Execution
    with concurrent.futures.ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        future_to_table = {executor.submit(verify_table, t): t for t in table_list}
        
        for future in concurrent.futures.as_completed(future_to_table):
            res = future.result()
            results_summary.append({
                "Table Name": res['table'],
                "Status": res['status'],
                "Diff Count": res.get('diff_count', 0),
                "Rows (Oracle)": res.get('row_ora', ''),
                "Rows (PolarDB)": res.get('row_plr', ''),
                "Notes": res.get('error', '')
            })
            
            if res['status'] == "MISMATCH":
                mismatch_details[res['table']] = res['data']

    # 5. Output to Excel
    with pd.ExcelWriter(OUTPUT_EXCEL, engine='openpyxl') as writer:
        # Summary Sheet
        summary_df = pd.DataFrame(results_summary)
        summary_df.to_excel(writer, sheet_name='Summary', index=False)

        # Apply summary formatting
        ws_summary = writer.sheets['Summary']
        header_fill = PatternFill("solid", fgColor="4472C4")
        header_font = Font(bold=True, color="FFFFFF")
        match_fill    = PatternFill("solid", fgColor="C6EFCE")  # green
        mismatch_fill = PatternFill("solid", fgColor="FCE4D6")  # orange
        error_fill    = PatternFill("solid", fgColor="FFCCCC")  # red

        for cell in ws_summary[1]:
            cell.fill = header_fill
            cell.font = header_font

        status_col = summary_df.columns.get_loc("Status") + 1
        for row in ws_summary.iter_rows(min_row=2):
            status_cell = row[status_col - 1]
            if status_cell.value == "MATCH":
                status_cell.fill = match_fill
            elif status_cell.value == "MISMATCH":
                status_cell.fill = mismatch_fill
            elif status_cell.value == "ERROR":
                status_cell.fill = error_fill

        # Detail Sheets for mismatches with formatting
        for table_name, diff_df in mismatch_details.items():
            sheet_name = f"Diff_{table_name}"[:31]
            diff_df.to_excel(writer, sheet_name=sheet_name, index=False)
            apply_diff_sheet_formatting(writer.sheets[sheet_name], diff_df)

    print(f"\n[!] Verification complete. Results saved to {OUTPUT_EXCEL}")

if __name__ == "__main__":
    main()